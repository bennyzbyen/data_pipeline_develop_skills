#!/usr/bin/env python3
"""Validate a DataHub Pipeline Export Excel workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DATA_UTILIZATION_SHEET = "Data Utilization"
TARGET_SHEET = "Target & Catalog"
FIELD_SHEET = "Target Field"
PIPELINE_SHEET = "Pipeline"

EXPECTED_HEADERS = {
    DATA_UTILIZATION_SHEET: ["*data_utilization_name", "data_utilization_description"],
    TARGET_SHEET: [
        "*data_utilization_name",
        "*target_name",
        "target_description",
        "data_storage_type",
        "db",
        "table_name",
        "table_description",
        "dataset_name",
        "dataset_title",
        "dataset_description",
        "dataset_business_owner",
        "dataset_business_owner_email",
        "dataset_it_owner",
        "dataset_it_owner_email",
        "dataset_fe",
        "dataset_fe_email",
        "dataset_it_bp",
        "dataset_it_bp_email",
        "dataset_data_engineer",
        "dataset_data_engineer_email",
        "dataset_source",
    ],
    FIELD_SHEET: [
        "*target_name",
        "*field_name",
        "*field_label",
        "field_description",
        "*field_type",
        "field_length",
        "field_sequence",
    ],
    PIPELINE_SHEET: [
        "*data_utilization_name",
        "*pipeline_name",
        "pipeline_description",
        "*enable",
        "*is_octopus",
        "pipeline_trigger",
        "pipeline_trigger_start",
        "pipeline_trigger_end",
        "pipeline_status_notification",
        "pipeline_notification_emails",
        "task1_name",
        "task1_description",
        "task1_link_target_names",
        "task1_mlp_params",
    ],
}

ALLOWED_STORAGE_TYPES = {"", "HBASE", "HDFS", "CLICKHOUSE", "SV_CLICKHOUSE", "MSSQL", "MYSQL"}
REQUIRED_FIELD_TYPE = "TEXT"


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def headers(ws, count: int) -> list[str]:
    return [clean(ws.cell(2, col).value) for col in range(1, count + 1)]


def row_dicts(ws, expected_headers: list[str]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for row_idx in range(3, ws.max_row + 1):
        row = {header: clean(ws.cell(row_idx, col).value) for col, header in enumerate(expected_headers, start=1)}
        if any(row.values()):
            row["_row"] = str(row_idx)
            rows.append(row)
    return rows


def required_headers(expected_headers: list[str]) -> list[str]:
    return [header for header in expected_headers if header.startswith("*")]


def add_duplicate_errors(values: list[str], label: str, errors: list[str]) -> None:
    seen: set[str] = set()
    duplicates: set[str] = set()
    for value in values:
        if not value:
            continue
        if value in seen:
            duplicates.add(value)
        seen.add(value)
    for value in sorted(duplicates):
        errors.append(f"Duplicate {label}: {value}")


def add_empty_inline_string_errors(ws, max_col: int, errors: list[str]) -> None:
    for (row_idx, col_idx), cell in ws._cells.items():
        if row_idx < 3 or col_idx > max_col:
            continue
        if cell.value is None and cell.data_type == "inlineStr":
            errors.append(f"{ws.title} {cell.coordinate} is an empty inline string cell; leave blank cells absent/null.")


def validate_workbook(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=False)
    errors: list[str] = []
    warnings: list[str] = []
    row_counts: dict[str, int] = {}
    data: dict[str, list[dict[str, str]]] = {}

    for sheet, expected in EXPECTED_HEADERS.items():
        if sheet not in wb.sheetnames:
            errors.append(f"Missing sheet: {sheet}")
            continue
        ws = wb[sheet]
        actual = headers(ws, len(expected))
        if actual != expected:
            errors.append(f"{sheet} headers mismatch: expected={expected!r} actual={actual!r}")
            continue
        rows = row_dicts(ws, expected)
        data[sheet] = rows
        row_counts[sheet] = len(rows)
        add_empty_inline_string_errors(ws, len(expected), errors)
        for row in rows:
            for header in required_headers(expected):
                if not row.get(header):
                    errors.append(f"{sheet} row {row['_row']} missing required column {header}")

    dus = data.get(DATA_UTILIZATION_SHEET, [])
    targets = data.get(TARGET_SHEET, [])
    fields = data.get(FIELD_SHEET, [])
    pipelines = data.get(PIPELINE_SHEET, [])

    du_names = {row["*data_utilization_name"] for row in dus if row.get("*data_utilization_name")}
    target_names = {row["*target_name"] for row in targets if row.get("*target_name")}

    add_duplicate_errors([row.get("*data_utilization_name", "") for row in dus], "data_utilization_name", errors)
    add_duplicate_errors([row.get("*target_name", "") for row in targets], "target_name", errors)
    add_duplicate_errors([row.get("*pipeline_name", "") for row in pipelines], "pipeline_name", errors)

    for row in targets:
        if row.get("*data_utilization_name") not in du_names:
            errors.append(f"Target & Catalog row {row['_row']} references unknown data_utilization_name: {row.get('*data_utilization_name')}")
        storage = row.get("data_storage_type", "")
        if storage not in ALLOWED_STORAGE_TYPES:
            errors.append(f"Target & Catalog row {row['_row']} has unsupported data_storage_type: {storage}")
        if not row.get("dataset_business_owner_email"):
            warnings.append(f"Target & Catalog row {row['_row']} dataset owner/email fields are blank.")
        for name_column in [
            "dataset_business_owner",
            "dataset_it_owner",
            "dataset_fe",
            "dataset_it_bp",
            "dataset_data_engineer",
        ]:
            name_value = row.get(name_column, "")
            if "@" in name_value or any(ch.isdigit() for ch in name_value):
                warnings.append(f"Target & Catalog row {row['_row']} {name_column} should be a person-style name without email or digits.")

    for row in fields:
        target = row.get("*target_name")
        if target not in target_names:
            errors.append(f"Target Field row {row['_row']} references unknown target_name: {target}")
        raw_field_type = row.get("*field_type")
        if clean(raw_field_type) != REQUIRED_FIELD_TYPE:
            errors.append(f"Target Field row {row['_row']} field_type must be TEXT, got: {raw_field_type}")

    field_ws = wb[FIELD_SHEET]
    sequence_col = EXPECTED_HEADERS[FIELD_SHEET].index("field_sequence") + 1
    for row_idx in range(3, field_ws.max_row + 1):
        cell = field_ws.cell(row_idx, sequence_col)
        if cell.value in (None, ""):
            continue
        if cell.data_type != "s":
            errors.append(f"Target Field row {row_idx} field_sequence must be stored as text, got cell type: {cell.data_type}")

    target_field_counts: dict[str, int] = {}
    for row in fields:
        target = row.get("*target_name", "")
        target_field_counts[target] = target_field_counts.get(target, 0) + 1
    for target in sorted(target_names):
        if target_field_counts.get(target, 0) == 0:
            warnings.append(f"Target has no field rows: {target}")

    for row in pipelines:
        if row.get("*data_utilization_name") not in du_names:
            errors.append(f"Pipeline row {row['_row']} references unknown data_utilization_name: {row.get('*data_utilization_name')}")
        if row.get("*enable") not in {"0", "1"}:
            errors.append(f"Pipeline row {row['_row']} has unsupported enable value: {row.get('*enable')}")
        if row.get("*is_octopus") not in {"0", "1"}:
            errors.append(f"Pipeline row {row['_row']} has unsupported is_octopus value: {row.get('*is_octopus')}")
        if not row.get("task1_link_target_names"):
            warnings.append(f"Pipeline row {row['_row']} task1_link_target_names is blank.")
        if not row.get("task1_mlp_params"):
            warnings.append(f"Pipeline row {row['_row']} task1_mlp_params is blank.")

    return {
        "ok": not errors,
        "workbook": str(path),
        "row_counts": row_counts,
        "errors": errors,
        "warnings": warnings,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Validate a Pipeline Export Excel workbook.")
    parser.add_argument("--xlsx", required=True, type=Path)
    parser.add_argument("--json-out", type=Path)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    result = validate_workbook(args.xlsx)
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(text + "\n", encoding="utf-8")
    print(text)
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
