#!/usr/bin/env python3
"""Build a DataHub Pipeline Export Excel workbook from structured facts."""

from __future__ import annotations

import argparse
import json
import re
from collections import OrderedDict, defaultdict
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DATA_UTILIZATION_SHEET = "Data Utilization"
TARGET_SHEET = "Target & Catalog"
FIELD_SHEET = "Target Field"
PIPELINE_SHEET = "Pipeline"

EXPECTED_HEADERS = {
    DATA_UTILIZATION_SHEET: ["*data_utilization_name", "data_utilization_description"],
    TARGET_SHEET: [
        "*data_utilization_name",
        "*target_name",
        "target_description",
        "data_storage_type",
        "db",
        "table_name",
        "table_description",
        "dataset_name",
        "dataset_title",
        "dataset_description",
        "dataset_business_owner",
        "dataset_business_owner_email",
        "dataset_it_owner",
        "dataset_it_owner_email",
        "dataset_fe",
        "dataset_fe_email",
        "dataset_it_bp",
        "dataset_it_bp_email",
        "dataset_data_engineer",
        "dataset_data_engineer_email",
        "dataset_source",
    ],
    FIELD_SHEET: [
        "*target_name",
        "*field_name",
        "*field_label",
        "field_description",
        "*field_type",
        "field_length",
        "field_sequence",
    ],
    PIPELINE_SHEET: [
        "*data_utilization_name",
        "*pipeline_name",
        "pipeline_description",
        "*enable",
        "*is_octopus",
        "pipeline_trigger",
        "pipeline_trigger_start",
        "pipeline_trigger_end",
        "pipeline_status_notification",
        "pipeline_notification_emails",
        "task1_name",
        "task1_description",
        "task1_link_target_names",
        "task1_mlp_params",
    ],
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def compact_space(value: str) -> str:
    return re.sub(r"\s+", " ", clean(value))


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def row_headers(ws, count: int) -> list[str]:
    return [clean(ws.cell(2, col).value) for col in range(1, count + 1)]


def ensure_template(wb) -> None:
    missing = [sheet for sheet in EXPECTED_HEADERS if sheet not in wb.sheetnames]
    if missing:
        raise ValueError(f"Template is missing required sheets: {missing}")
    for sheet, expected in EXPECTED_HEADERS.items():
        actual = row_headers(wb[sheet], len(expected))
        if actual != expected:
            raise ValueError(f"{sheet} headers mismatch. expected={expected!r} actual={actual!r}")


def copy_row_style(ws, source_row: int, target_row: int, max_col: int) -> None:
    if source_row == target_row:
        return
    if source_row in ws.row_dimensions:
        ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.protection:
            dst.protection = copy(src.protection)


def prepare_data_area(ws, rows_needed: int, max_col: int) -> None:
    end_row = max(3, 2 + rows_needed)
    if ws.max_row < end_row:
        ws.insert_rows(ws.max_row + 1, amount=end_row - ws.max_row)
    style_row = 3 if ws.max_row >= 3 else 2
    for row in range(3, end_row + 1):
        copy_row_style(ws, style_row, row, max_col)
    for row in range(3, ws.max_row + 1):
        for col in range(1, max_col + 1):
            ws.cell(row, col).value = None


def write_rows(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    prepare_data_area(ws, len(rows), len(headers))
    for offset, row in enumerate(rows, start=3):
        for col, header in enumerate(headers, start=1):
            ws.cell(offset, col).value = row.get(header, "")


def ordered_unique(values: list[str]) -> list[str]:
    result: OrderedDict[str, None] = OrderedDict()
    for value in values:
        value = clean(value)
        if value:
            result[value] = None
    return list(result.keys())


def storage_type(*values: str) -> str:
    text = " ".join(clean(v).lower() for v in values)
    if "superview" in text or "sv_clickhouse" in text or "sv clickhouse" in text:
        return "SV_CLICKHOUSE"
    if "clickhouse" in text:
        return "CLICKHOUSE"
    if "hbase" in text:
        return "HBASE"
    if "hdfs" in text or "fs" in text:
        return "HDFS"
    if "mssql" in text or "sql server" in text:
        return "MSSQL"
    if "mysql" in text:
        return "MYSQL"
    return clean(values[0]) if values else ""


def field_type(value: str) -> str:
    key = clean(value).upper()
    key = key.replace("DATATIME", "DATETIME")
    if key.startswith("VARCHAR") or key.startswith("CHAR") or key in {"STRING", "TEXT"}:
        return "TEXT"
    if key in {"INT", "INTEGER"}:
        return "INT"
    if key in {"LONG", "BIGINT"}:
        return "LONG"
    if key.startswith("DECIMAL") or key in {"DOUBLE", "FLOAT", "NUMBER"}:
        return "DECIMAL"
    if key in {"DATE", "DATA"}:
        return "DATE"
    if key in {"DATETIME", "TIMESTAMP"}:
        return "TIMESTAMP"
    if key in {"BOOLEAN", "BOOL"}:
        return "BOOLEAN"
    return clean(value)


def split_physical_table(table: str, database: str, storage: str) -> tuple[str, str]:
    table = clean(table)
    database = clean(database)
    mapped_storage = storage_type(storage)
    if mapped_storage == "HBASE":
        return "", table
    if database and table.lower().startswith(database.lower() + "."):
        return database, table[len(database) + 1 :]
    if "." in table and not database:
        left, right = table.split(".", 1)
        return left, right
    return database, table


def normalize_table_key(value: str) -> str:
    value = clean(value).lower()
    if "." in value:
        return value.split(".")[-1]
    return value


def extract_email(value: str) -> str:
    match = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", clean(value))
    return match.group(0) if match else ""


def strip_digits(value: str) -> str:
    return re.sub(r"\d+", "", clean(value)).strip()


def contact_name(value: str) -> str:
    value = clean(value)
    email = extract_email(value)
    if email:
        non_email = clean(value.replace(email, ""))
        non_email = re.sub(r"[<>()；;，,]+", " ", non_email).strip()
        if non_email:
            return strip_digits(non_email)
        local = email.split("@", 1)[0]
        local = strip_digits(local)
        parts = [part for part in re.split(r"[._\-\s]+", local) if part]
        return " ".join(part[:1].upper() + part[1:] for part in parts)
    return strip_digits(value)


def data_utilization_descriptions(facts: dict[str, Any]) -> dict[str, str]:
    descriptions: dict[str, str] = {}
    for item in facts.get("data_utilizations", []):
        name = clean(item.get("name"))
        description = clean(item.get("description"))
        if name and description and name not in descriptions:
            descriptions[name] = description
    return descriptions


def catalog_info_by_table(facts: dict[str, Any]) -> dict[str, dict[str, Any]]:
    catalog: dict[str, dict[str, Any]] = {}
    for item in facts.get("report_catalog_basic_info", []):
        key = normalize_table_key(clean(item.get("data_item")))
        if key and key not in catalog:
            catalog[key] = item
    return catalog


def find_physical(target: dict[str, Any], physical_targets: list[dict[str, Any]]) -> dict[str, Any]:
    physical_table = clean(target.get("physical_table"))
    physical_desc = clean(target.get("physical_description") or target.get("description"))
    target_table_key = normalize_table_key(physical_table)
    target_desc = compact_space(physical_desc).lower()
    for item in physical_targets:
        if clean(item.get("table")) and clean(item.get("table")).lower() == physical_table.lower():
            return item
    for item in physical_targets:
        if normalize_table_key(clean(item.get("table") or item.get("table_name"))) == target_table_key:
            return item
    for item in physical_targets:
        if target_desc and compact_space(item.get("description", "")).lower() == target_desc:
            return item
    return {}


def build_data_utilization_rows(facts: dict[str, Any], override_name: str, project_name: str) -> list[dict[str, Any]]:
    if override_name:
        names = [override_name]
    else:
        names = ordered_unique([clean(item.get("data_utilization")) for item in facts.get("report_targets", [])])
    if not names and project_name:
        names = [project_name.lower()]
    descriptions = data_utilization_descriptions(facts)
    rows = []
    for name in names:
        rows.append(
            {
                "*data_utilization_name": name,
                "data_utilization_description": descriptions.get(name, ""),
            }
        )
    return rows


def build_target_rows(
    facts: dict[str, Any],
    data_utilization: str,
    questions: list[str],
) -> list[dict[str, Any]]:
    physical_targets = facts.get("report_physical_targets", [])
    catalog_by_table = catalog_info_by_table(facts)
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    missing_catalog: list[str] = []
    for item in facts.get("report_targets", []):
        target_name = clean(item.get("target_name"))
        if not target_name or target_name in seen:
            continue
        seen.add(target_name)
        physical = find_physical(item, physical_targets)
        raw_storage = clean(physical.get("storage") or item.get("storage"))
        mapped_storage = storage_type(raw_storage)
        table = clean(physical.get("table") or item.get("physical_table") or item.get("table_name"))
        database = clean(physical.get("database") or item.get("database"))
        db, table_name = split_physical_table(table, database, mapped_storage)
        description = clean(item.get("description") or physical.get("description"))
        catalog = catalog_by_table.get(normalize_table_key(table_name)) or catalog_by_table.get(normalize_table_key(table))
        if not catalog:
            missing_catalog.append(table_name or target_name)
        business_owner_value = clean(catalog.get("business_owner")) if catalog else ""
        it_owner_value = clean(catalog.get("it_owner")) if catalog else ""
        fe_value = clean(catalog.get("fe")) if catalog else ""
        it_bp_value = clean(catalog.get("it_bp")) if catalog else ""
        data_engineer_value = clean(catalog.get("data_engineer")) if catalog else ""
        catalog_title = clean(catalog.get("title")) if catalog else ""
        catalog_data_item = clean(catalog.get("data_item")) if catalog else ""
        rows.append(
            {
                "*data_utilization_name": data_utilization or clean(item.get("data_utilization")),
                "*target_name": target_name,
                "target_description": description,
                "data_storage_type": mapped_storage,
                "db": db,
                "table_name": table_name,
                "table_description": clean(physical.get("description") or item.get("physical_description") or description),
                "dataset_name": catalog_data_item,
                "dataset_title": catalog_title,
                "dataset_description": catalog_title,
                "dataset_business_owner": contact_name(business_owner_value),
                "dataset_business_owner_email": extract_email(business_owner_value),
                "dataset_it_owner": contact_name(it_owner_value),
                "dataset_it_owner_email": extract_email(it_owner_value),
                "dataset_fe": contact_name(fe_value),
                "dataset_fe_email": extract_email(fe_value),
                "dataset_it_bp": contact_name(it_bp_value),
                "dataset_it_bp_email": extract_email(it_bp_value),
                "dataset_data_engineer": contact_name(data_engineer_value),
                "dataset_data_engineer_email": extract_email(data_engineer_value),
                "dataset_source": "",
            }
        )
    for item in missing_catalog:
        questions.append(f"Target & Catalog: no Catalog Basic Info row matched `{item}`.")
    return rows


def mapping_target_name(item: dict[str, Any]) -> str:
    return clean(item.get("inferred_target_name") or item.get("target_name") or item.get("declared_target_name"))


def mapping_priority(item: dict[str, Any]) -> tuple[int, int]:
    source_index = item.get("source_doc_index")
    if not isinstance(source_index, int):
        source_index = 99
    fields = item.get("fields", [])
    field_count_score = -len(fields) if isinstance(fields, list) else 0
    return source_index, field_count_score


def choose_field_mappings(facts: dict[str, Any], target_names: list[str]) -> tuple[dict[str, dict[str, Any]], list[str]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in facts.get("report_field_mappings", []):
        target = mapping_target_name(item)
        if target in target_names:
            grouped[target].append(item)
    chosen: dict[str, dict[str, Any]] = {}
    missing: list[str] = []
    for target in target_names:
        candidates = grouped.get(target, [])
        if not candidates:
            missing.append(target)
            continue
        chosen[target] = sorted(candidates, key=mapping_priority)[0]
    return chosen, missing


def field_key(field: dict[str, Any]) -> str:
    return clean(
        field.get("target_field")
        or field.get("field_key")
        or field.get("field_name_en")
        or field.get("column")
        or field.get("Column")
    )


def field_label(field: dict[str, Any], fallback: str) -> str:
    return clean(field.get("field_name") or field.get("field_label") or field.get("description") or fallback)


def field_description(field: dict[str, Any]) -> str:
    return clean(field.get("field_description") or field.get("calculation_logic") or field.get("remark"))


def build_field_rows(
    facts: dict[str, Any],
    target_names: list[str],
    questions: list[str],
) -> list[dict[str, Any]]:
    chosen, missing = choose_field_mappings(facts, target_names)
    for target in missing:
        questions.append(f"Target Field: target `{target}` has no high-confidence field dictionary.")
    rows: list[dict[str, Any]] = []
    for target in target_names:
        mapping = chosen.get(target)
        if not mapping:
            continue
        sequence = 0
        seen_fields: set[str] = set()
        for field in mapping.get("fields", []):
            key = field_key(field)
            if not key or key in seen_fields:
                continue
            seen_fields.add(key)
            mapped_type = field_type(clean(field.get("field_type") or field.get("type") or field.get("Type")))
            if not mapped_type:
                questions.append(f"Target Field: `{target}.{key}` has no confirmed field type.")
            rows.append(
                {
                    "*target_name": target,
                    "*field_name": key,
                    "*field_label": field_label(field, key),
                    "field_description": field_description(field),
                    "*field_type": mapped_type,
                    "field_length": "",
                    "field_sequence": sequence,
                }
            )
            sequence += 1
    return rows


def build_pipeline_rows(
    facts: dict[str, Any],
    data_utilization: str,
    questions: list[str],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in facts.get("report_schedules", []):
        pipeline_name = clean(item.get("pipeline_name"))
        if not pipeline_name or pipeline_name in seen:
            continue
        seen.add(pipeline_name)
        task_name = clean(item.get("task_name"))
        rows.append(
            {
                "*data_utilization_name": data_utilization or clean(item.get("data_utilization")),
                "*pipeline_name": pipeline_name,
                "pipeline_description": clean(item.get("description")),
                "*enable": "1",
                "*is_octopus": "0",
                "pipeline_trigger": "",
                "pipeline_trigger_start": "",
                "pipeline_trigger_end": "",
                "pipeline_status_notification": "FINISHED,FAILED",
                "pipeline_notification_emails": "",
                "task1_name": task_name,
                "task1_description": clean(item.get("description")),
                "task1_link_target_names": "",
                "task1_mlp_params": "",
            }
        )
        questions.append(f"Pipeline: `{pipeline_name}` task-target links and MLP params require manual confirmation.")
    return rows


def walk_strings(value: Any) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, list):
        result: list[str] = []
        for item in value:
            result.extend(walk_strings(item))
        return result
    if isinstance(value, dict):
        result = []
        for item in value.values():
            result.extend(walk_strings(item))
        return result
    return []


def add_conflict_questions(facts: dict[str, Any], target_count: int, questions: list[str]) -> None:
    for conflict in facts.get("conflicts", []):
        if isinstance(conflict, dict):
            kind = clean(conflict.get("kind"))
            target = clean(conflict.get("target"))
            field = clean(conflict.get("field"))
            questions.append(f"Conflict: {kind} target={target} field={field}; review source documents before import.")
    claims: set[int] = set()
    for text in walk_strings(facts):
        for match in re.finditer(r"(\d+)\s*个\s*(?:target|Target|TARGET)", text):
            claims.add(int(match.group(1)))
    for claim in sorted(claims):
        if claim != target_count:
            questions.append(f"Target count conflict: document text mentions {claim} targets, generated workbook has {target_count}.")


def write_questions(path: Path, project_name: str, questions: list[str], counts: dict[str, int]) -> None:
    unique_questions = ordered_unique(questions)
    lines = [
        f"# {project_name or 'Pipeline Excel'} Questions",
        "",
        "## Generated Row Counts",
        "",
        "| Sheet | Rows |",
        "| --- | ---: |",
    ]
    for sheet, count in counts.items():
        lines.append(f"| {sheet} | {count} |")
    lines.extend(["", "## To Confirm", ""])
    if unique_questions:
        for item in unique_questions:
            lines.append(f"- {item}")
    else:
        lines.append("- No open questions were generated.")
    lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def build_workbook(args: argparse.Namespace) -> dict[str, int]:
    facts = read_json(args.structured_facts)
    wb = load_workbook(args.template_xlsx)
    ensure_template(wb)

    questions: list[str] = []
    du_rows = build_data_utilization_rows(facts, args.data_utilization, args.project_name)
    selected_du = args.data_utilization or (du_rows[0]["*data_utilization_name"] if du_rows else "")
    target_rows = build_target_rows(facts, selected_du, questions)
    target_names = [row["*target_name"] for row in target_rows]
    field_rows = build_field_rows(facts, target_names, questions)
    pipeline_rows = build_pipeline_rows(facts, selected_du, questions)
    add_conflict_questions(facts, len(target_rows), questions)

    write_rows(wb[DATA_UTILIZATION_SHEET], EXPECTED_HEADERS[DATA_UTILIZATION_SHEET], du_rows)
    write_rows(wb[TARGET_SHEET], EXPECTED_HEADERS[TARGET_SHEET], target_rows)
    write_rows(wb[FIELD_SHEET], EXPECTED_HEADERS[FIELD_SHEET], field_rows)
    write_rows(wb[PIPELINE_SHEET], EXPECTED_HEADERS[PIPELINE_SHEET], pipeline_rows)

    args.out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out_xlsx)

    counts = {
        DATA_UTILIZATION_SHEET: len(du_rows),
        TARGET_SHEET: len(target_rows),
        FIELD_SHEET: len(field_rows),
        PIPELINE_SHEET: len(pipeline_rows),
    }
    write_questions(args.questions_out, args.project_name, questions, counts)
    return counts


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build a Pipeline Export Excel workbook from structured facts.")
    parser.add_argument("--template-xlsx", required=True, type=Path)
    parser.add_argument("--structured-facts", required=True, type=Path)
    parser.add_argument("--out-xlsx", required=True, type=Path)
    parser.add_argument("--questions-out", required=True, type=Path)
    parser.add_argument("--project-name", required=True)
    parser.add_argument("--data-utilization", default="")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    counts = build_workbook(args)
    print(json.dumps({"out_xlsx": str(args.out_xlsx), "questions": str(args.questions_out), "row_counts": counts}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
